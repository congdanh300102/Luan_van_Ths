import sys
import os
from docx import Document
from docx.shared import Pt

def add_paragraph(doc, line):
    # Simple heuristic: headings start with '#'
    stripped = line.lstrip()
    if stripped.startswith('### '):
        p = doc.add_heading(stripped[4:], level=3)
    elif stripped.startswith('## '):
        p = doc.add_heading(stripped[3:], level=2)
    elif stripped.startswith('# '):
        p = doc.add_heading(stripped[2:], level=1)
    elif stripped.startswith('|') and '|' in stripped[1:]:
        # Table row – we'll just add as plain text for simplicity
        p = doc.add_paragraph(line)
    else:
        p = doc.add_paragraph(line)
    return p

def main(md_path, docx_path):
    if not os.path.isfile(md_path):
        print(f"Markdown file not found: {md_path}")
        sys.exit(1)
    doc = Document()
    with open(md_path, encoding='utf-8') as f:
        for line in f:
            line = line.rstrip('\n')
            if line == '':
                doc.add_paragraph('')
            else:
                add_paragraph(doc, line)
    doc.save(docx_path)
    print(f"Created docx: {docx_path}")

if __name__ == '__main__':
    if len(sys.argv) != 3:
        print('Usage: python generate_docx.py <input.md> <output.docx>')
        sys.exit(1)
    main(sys.argv[1], sys.argv[2])

